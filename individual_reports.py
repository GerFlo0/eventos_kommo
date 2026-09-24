"""
Reporte de estado actual e historial de leads de Kommo.

Uso desde otro script:

    from individual_reports import generar_reporte_estado_leads

    generar_reporte_estado_leads(resultado_df, asesor="NOMBRE DEL ASESOR")

`resultado_df` es el resultado de la consulta: el historial de etapas de los
leads del asesor (filas de historial_etapas_kommo.xlsx). También acepta la
ruta a un .xlsx.

El reporte incluye tres hojas: Historial por lead, Resumen y Estado actual.
"""

import unicodedata
import re
from datetime import datetime
from pathlib import Path

import duckdb as db
import pandas as pd

import functions as fn
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

RUTA_HISTORIAL = "tablas/historial_etapas_kommo.xlsx"
CARPETA_REPORTES = "tablas/reportes/individuales"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# Columnas mínimas que debe traer el resultado de la consulta.
COLUMNAS_REQUERIDAS = ["LEAD_ID", "ETAPA_ANTERIOR", "ETAPA_NUEVA", "FECHA"]
# Columnas opcionales: si no vienen, se crean vacías.
COLUMNAS_OPCIONALES = ["DIAS_EN_ETAPA_ANTERIOR", "ETIQUETAS"]

# Categorías de estado actual, en el orden en que aparecen en el reporte.
# Cada categoría se detecta comparando ETAPA_NUEVA sin mayúsculas ni acentos
# contra la lista de nombres equivalentes.
CATEGORIAS = [
    ("Logrado con éxito", ["logrado con exito", "otorgado"]),
    ("Oferta", ["oferta"]),
    ("Oferta en espera", ["oferta en espera"]),
    ("Documentación", ["documentacion"]),
    ("Capturado", ["capturado"]),
    ("Venta perdida", ["venta perdida", "ventas perdidas", "ventas perdidos",
                       "lead perdido", "leads perdidos"]),
    ("Sin capacidad", ["sin capacidad"]),
]
CATEGORIA_FUTURO = "Dejado a futuro"   # etapas cuyo nombre contiene un año (2025, 2026...)
# Un año de 4 dígitos que empieza con 20 y no forma parte de un número más
# largo. Antes bastaba con contener "20", y etapas como "OFERTA 20%" o
# "Llamar en 20 días" se clasificaban por error como "Dejado a futuro".
PATRON_FUTURO = re.compile(r"(?<!\d)20\d{2}(?!\d)")
CATEGORIA_OTROS = "Otros"              # cualquier etapa no contemplada arriba
ORDEN_CATEGORIAS = [c for c, _ in CATEGORIAS] + [CATEGORIA_FUTURO, CATEGORIA_OTROS]

_MAPA_ETAPAS = {alias: cat for cat, aliases in CATEGORIAS for alias in aliases}

FORMATO_FECHA = "dd/mm/yyyy hh:mm"
FUENTE = "Arial"


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def normalizar(texto):
    """Minúsculas, sin acentos y con espacios simples."""
    if texto is None or pd.isna(texto):
        return ""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().split())


def clasificar_etapa(etapa):
    """Devuelve la categoría de estado actual para una etapa de Kommo."""
    etapa_norm = normalizar(etapa)
    if PATRON_FUTURO.search(etapa_norm):
        return CATEGORIA_FUTURO
    return _MAPA_ETAPAS.get(etapa_norm, CATEGORIA_OTROS)


def formato_tiempo(dias):
    """Convierte días decimales a texto legible, p. ej. 12.4 -> '12 d 10 h'."""
    if dias is None or pd.isna(dias):
        return ""
    total_horas = int(round(dias * 24))
    d, h = divmod(total_horas, 24)
    return f"{d} d {h} h"


def construir_ruta(grupo):
    """Ruta de etapas de un lead: 'Contactado → oferta (01/09/2026 09:20) → ...'."""
    primera_anterior = grupo["ETAPA_ANTERIOR"].iloc[0]
    partes = [str(primera_anterior) if pd.notna(primera_anterior) else "(sin etapa previa)"]
    for etapa, fecha in zip(grupo["ETAPA_NUEVA"], grupo["FECHA"]):
        partes.append(f"{etapa} ({fecha:%d/%m/%Y %H:%M})")
    return " → ".join(partes)


# ---------------------------------------------------------------------------
# Procesamiento
# ---------------------------------------------------------------------------
def filtrar_por_etiqueta(df, etiqueta):
    """
    Devuelve el historial completo de los leads que tienen la etiqueta en
    alguno de sus registros. La comparación no distingue mayúsculas ni acentos.
    """
    con = db.connect()
    con.register("historial", df)
    query = """
        SELECT * FROM historial
        WHERE LEAD_ID IN (
            SELECT LEAD_ID FROM historial
            WHERE strip_accents(ETIQUETAS) ILIKE strip_accents($1)
        )
    """
    resultado = con.execute(query, [f"%{etiqueta}%"]).df()
    con.close()
    return resultado


def preparar_historial(df, fecha_corte):
    """Ordena el historial y agrega paso, último movimiento y días en cada etapa."""
    con = db.connect()
    con.register("historial", df)
    hist = con.execute("""
        SELECT *,
            ROW_NUMBER() OVER w                 AS PASO,
            COUNT(*) OVER (PARTITION BY LEAD_ID) AS TOTAL_MOVIMIENTOS,
            LEAD(FECHA) OVER w                  AS FECHA_SIGUIENTE
        FROM historial
        WINDOW w AS (PARTITION BY LEAD_ID ORDER BY FECHA)
        ORDER BY LEAD_ID, FECHA
    """).df()
    con.close()

    hist["FECHA"] = pd.to_datetime(hist["FECHA"])
    hist["FECHA_SIGUIENTE"] = pd.to_datetime(hist["FECHA_SIGUIENTE"])
    hist["ES_ULTIMO_MOVIMIENTO"] = hist["PASO"] == hist["TOTAL_MOVIMIENTOS"]

    # Días en la etapa nueva: hasta el siguiente movimiento, o hasta la fecha
    # de corte si es la etapa actual.
    fin_etapa = hist["FECHA_SIGUIENTE"].fillna(pd.Timestamp(fecha_corte))
    hist["DIAS_EN_ETAPA_NUEVA"] = ((fin_etapa - hist["FECHA"]).dt.total_seconds() / 86400).round(2)
    return hist


def generar_reporte(df, fecha_corte):
    """Devuelve un dict con los DataFrames de cada hoja del reporte."""
    df = df.copy()
    df["FECHA"] = pd.to_datetime(df["FECHA"])
    hist = preparar_historial(df, fecha_corte)
    for col in COLUMNAS_OPCIONALES:
        if col not in hist.columns:
            hist[col] = None

    # --- Estado actual: último movimiento de cada lead -----------------------
    ultimos = hist[hist["ES_ULTIMO_MOVIMIENTO"]].copy()
    ultimos["ESTADO_ACTUAL"] = ultimos["ETAPA_NUEVA"].map(clasificar_etapa)
    ultimos["DIAS_DESDE_ULTIMO_MOVIMIENTO"] = (
        (pd.Timestamp(fecha_corte) - ultimos["FECHA"]).dt.total_seconds() / 86400
    ).round(2)
    ultimos["TIEMPO_TRANSCURRIDO"] = ultimos["DIAS_DESDE_ULTIMO_MOVIMIENTO"].map(formato_tiempo)

    # Las etiquetas pueden venir solo en algunos registros: se toma la última no vacía.
    etiquetas = hist.groupby("LEAD_ID")["ETIQUETAS"].last()
    rutas = hist.groupby("LEAD_ID")[["ETAPA_ANTERIOR", "ETAPA_NUEVA", "FECHA"]].apply(construir_ruta)
    ultimos["ETIQUETAS"] = ultimos["LEAD_ID"].map(etiquetas)
    ultimos["RUTA"] = ultimos["LEAD_ID"].map(rutas)

    ultimos["_orden"] = ultimos["ESTADO_ACTUAL"].map(ORDEN_CATEGORIAS.index)
    ultimos = ultimos.sort_values(["_orden", "DIAS_DESDE_ULTIMO_MOVIMIENTO"], ascending=[True, False])

    columnas_estado = [
        "LEAD_ID", "creacion_de_lead","ESTADO_ACTUAL", "ETAPA_NUEVA", "FECHA",
        "DIAS_DESDE_ULTIMO_MOVIMIENTO", "TIEMPO_TRANSCURRIDO", "ETAPA_ANTERIOR",
        "ESTATUS DE NEGOCIO", "FECHA DICTAMEN", "TOTAL_MOVIMIENTOS", "RUTA",
        "ETIQUETAS", "EMBUDO",
    ]
    estado = (
        ultimos[[c for c in columnas_estado if c in ultimos.columns]]
        .rename(columns={"ETAPA_NUEVA": "ETAPA_ACTUAL", "FECHA": "FECHA_ULTIMO_MOVIMIENTO"})
        .reset_index(drop=True)
    )

    # --- Resumen por categoría ----------------------------------------------
    total = len(estado)
    filas = []
    for cat in ORDEN_CATEGORIAS:
        sub = estado[estado["ESTADO_ACTUAL"] == cat]
        filas.append({
            "ESTADO_ACTUAL": cat,
            "LEADS": len(sub),
            "PORCENTAJE": len(sub) / total if total else 0,
            "PROMEDIO_DIAS_DESDE_ULTIMO_MOVIMIENTO": round(sub["DIAS_DESDE_ULTIMO_MOVIMIENTO"].mean(), 2) if len(sub) else None,
        })
    filas.append({
        "ESTADO_ACTUAL": "TOTAL",
        "LEADS": total,
        "PORCENTAJE": 1 if total else 0,
        "PROMEDIO_DIAS_DESDE_ULTIMO_MOVIMIENTO": round(estado["DIAS_DESDE_ULTIMO_MOVIMIENTO"].mean(), 2) if total else None,
    })
    resumen = pd.DataFrame(filas)

    # --- Historial ------------------------------------------------------------
    hist["ESTADO_ACTUAL_LEAD"] = hist["LEAD_ID"].map(estado.set_index("LEAD_ID")["ESTADO_ACTUAL"])
    hist["ES_ULTIMO_MOVIMIENTO"] = hist["ES_ULTIMO_MOVIMIENTO"].map({True: "Sí", False: ""})
    historial = hist[[
        "LEAD_ID", "PASO", "ETAPA_ANTERIOR", "ETAPA_NUEVA", "FECHA",
        "DIAS_EN_ETAPA_ANTERIOR", "DIAS_EN_ETAPA_NUEVA",
        "ESTADO_ACTUAL_LEAD", "ES_ULTIMO_MOVIMIENTO",
    ]].reset_index(drop=True)

    return {"Resumen": resumen, "Estado actual": estado, "Historial": historial}


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
_RELLENO_ENCABEZADO = PatternFill("solid", start_color="1F3864")
_ANCHOS = {"RUTA": 90, "ETIQUETAS": 60}


def _dar_formato(ws, fila_encabezado, formatos=None):
    """Fuente, encabezado, anchos, filtros y paneles fijos para una hoja."""
    formatos = formatos or {}
    encabezados = [c.value for c in ws[fila_encabezado]]

    for fila in ws.iter_rows():
        for celda in fila:
            celda.font = Font(name=FUENTE, bold=celda.font.bold)

    for celda in ws[fila_encabezado]:
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, nombre in enumerate(encabezados, start=1):
        letra = ws.cell(row=fila_encabezado, column=idx).column_letter
        if nombre in formatos:
            for (celda,) in ws.iter_rows(min_row=fila_encabezado + 1, min_col=idx, max_col=idx):
                celda.number_format = formatos[nombre]
        largo = max(
            (len(str(c.value)) for (c,) in ws.iter_rows(min_row=fila_encabezado, min_col=idx, max_col=idx) if c.value is not None),
            default=10,
        )
        ws.column_dimensions[letra].width = _ANCHOS.get(nombre, min(max(largo + 2, 12), 40))

    ws.row_dimensions[fila_encabezado].height = 30
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=2)
    ws.auto_filter.ref = f"A{fila_encabezado}:{ws.cell(row=ws.max_row, column=len(encabezados)).coordinate}"


_RELLENO_LEAD = PatternFill("solid", start_color="D9E1F2")
_RELLENO_SUBENCABEZADO = PatternFill("solid", start_color="8EA9DB")
HOJA_POR_LEAD = "Historial por lead"
# Columnas que se muestran en la hoja "Estado actual" (el DataFrame interno
# conserva más datos porque los usan el Resumen y el Historial por lead).
COLUMNAS_HOJA_ESTADO = [
    "LEAD_ID", "creacion_de_lead","ETAPA_ACTUAL", "FECHA_ULTIMO_MOVIMIENTO",
    "TIEMPO_TRANSCURRIDO", "TOTAL_MOVIMIENTOS", "RUTA",
]
COLUMNAS_BLOQUE = [
    ("PASO", "PASO", 8, None),
    ("ETAPA_ANTERIOR", "ETAPA ANTERIOR", 22, None),
    ("ETAPA_NUEVA", "ETAPA NUEVA", 22, None),
    ("FECHA", "FECHA DEL CAMBIO", 20, FORMATO_FECHA),
    ("DIAS_EN_ETAPA_ANTERIOR", "DÍAS EN ETAPA ANTERIOR", 16, "0.00"),
    ("DIAS_EN_ETAPA_NUEVA", "DÍAS EN ETAPA NUEVA", 16, "0.00"),
]


def _escribir_historial_por_lead(libro, estado, historial):
    """
    Una hoja con un bloque por lead: encabezado con el lead, su etapa actual,
    la fecha de su última actualización de etapa y el tiempo transcurrido desde
    entonces; debajo, la tabla con todos sus cambios de etapa en orden.
    Devuelve {LEAD_ID: fila donde empieza su bloque} para crear vínculos.
    """
    ws = libro.create_sheet(HOJA_POR_LEAD, 0)
    n_cols = len(COLUMNAS_BLOQUE)
    negrita = Font(name=FUENTE, bold=True)
    normal = Font(name=FUENTE)
    inicio_bloques = {}
    fila = 1

    for _, lead in estado.iterrows():
        lead_id = lead["LEAD_ID"]
        inicio_bloques[lead_id] = fila

        # Fila 1 del bloque: identificación del lead
        titulo = f"LEAD {lead_id}  |  ETAPA ACTUAL: {str(lead['ETAPA_ACTUAL']).upper()}"
        ws.cell(row=fila, column=1, value=titulo).font = Font(name=FUENTE, bold=True, color="FFFFFF", size=12)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_cols)
        for col in range(1, n_cols + 1):
            ws.cell(row=fila, column=col).fill = _RELLENO_ENCABEZADO
        fila += 1

        # Fila 2: última actualización de etapa y tiempo transcurrido
        ws.cell(row=fila, column=1, value="Última actualización").font = negrita
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        c = ws.cell(row=fila, column=3, value=lead["FECHA_ULTIMO_MOVIMIENTO"].to_pydatetime())
        c.font = normal
        c.number_format = FORMATO_FECHA
        c.alignment = Alignment(horizontal="left")
        ws.cell(row=fila, column=4, value="Tiempo transcurrido").font = negrita
        ws.cell(
            row=fila, column=5,
            value=f"{lead['TIEMPO_TRANSCURRIDO']}  ({lead['DIAS_DESDE_ULTIMO_MOVIMIENTO']:.2f} días)",
        ).font = negrita
        ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=n_cols)
        for col in range(1, n_cols + 1):
            ws.cell(row=fila, column=col).fill = _RELLENO_LEAD
        fila += 1

        # Encabezado de la tabla de cambios
        for col, (_, titulo_col, _, _) in enumerate(COLUMNAS_BLOQUE, start=1):
            c = ws.cell(row=fila, column=col, value=titulo_col)
            c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
            c.fill = _RELLENO_SUBENCABEZADO
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        fila += 1

        # Cambios de etapa en orden cronológico; el más reciente en negritas
        movimientos = historial[historial["LEAD_ID"] == lead_id].sort_values("PASO")
        for _, mov in movimientos.iterrows():
            es_ultimo = mov["ES_ULTIMO_MOVIMIENTO"] == "Sí"
            for col, (campo, _, _, fmt) in enumerate(COLUMNAS_BLOQUE, start=1):
                valor = mov[campo]
                if pd.isna(valor):
                    valor = None
                elif campo == "FECHA":
                    valor = valor.to_pydatetime()
                c = ws.cell(row=fila, column=col, value=valor)
                c.font = negrita if es_ultimo else normal
                if fmt:
                    c.number_format = fmt
                if campo == "PASO":
                    c.alignment = Alignment(horizontal="center")
            fila += 1

        fila += 1  # fila en blanco entre leads

    for col, (_, _, ancho, _) in enumerate(COLUMNAS_BLOQUE, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.sheet_view.showGridLines = False
    return inicio_bloques


def escribir_excel(hojas, ruta, fecha_corte, asesor=""):
    """Escribe el reporte con fechas reales de Excel y formato legible."""
    with pd.ExcelWriter(ruta, engine="openpyxl", datetime_format=FORMATO_FECHA) as writer:
        fila_tabla_resumen = 5
        hojas["Resumen"].to_excel(writer, sheet_name="Resumen", index=False, startrow=fila_tabla_resumen - 1)
        estado = hojas["Estado actual"]
        # creacion_de_lead solo existe si el historial se generó con la versión
        # actual del extractor; si falta, la hoja sale sin esa columna.
        estado[[c for c in COLUMNAS_HOJA_ESTADO if c in estado.columns]].to_excel(
            writer, sheet_name="Estado actual", index=False)

        libro = writer.book

        # --- Historial por lead (primera hoja)
        inicio_bloques = _escribir_historial_por_lead(libro, hojas["Estado actual"], hojas["Historial"])
        libro.active = 0

        # --- Resumen
        ws = libro["Resumen"]
        ws["A1"], ws["B1"] = "Fecha de corte", fecha_corte
        ws["B1"].number_format = FORMATO_FECHA
        ws["A2"], ws["B2"] = "Asesor", asesor
        _dar_formato(ws, fila_tabla_resumen, {"PORCENTAJE": "0.0%", "PROMEDIO_DIAS_DESDE_ULTIMO_MOVIMIENTO": "0.00"})
        ws.auto_filter.ref = None
        ws.freeze_panes = None
        for c in ("A1", "A2"):
            ws[c].font = Font(name=FUENTE, bold=True)
        for celda in ws[ws.max_row]:  # fila TOTAL
            celda.font = Font(name=FUENTE, bold=True)
        nota = ws.max_row + 2
        ws.cell(row=nota, column=1, value="Notas:").font = Font(name=FUENTE, bold=True)
        notas = [
            "El tiempo transcurrido se calcula desde la última actualización de etapa hasta la fecha de corte (momento en que se generó el reporte).",
            "'Dejado a futuro' agrupa las etapas cuyo nombre contiene un año (p. ej. 2025 o 2026).",
            "'Otros' agrupa etapas no contempladas en la clasificación; revisarlas si aparecen.",
            "En 'Historial por lead', los días en la etapa nueva del último cambio se cuentan hasta la fecha de corte.",
        ]
        for i, texto in enumerate(notas, start=1):
            ws.cell(row=nota + i, column=1, value=texto).font = Font(name=FUENTE)
        ws.column_dimensions["A"].width = 22

        # --- Estado actual, con vínculo de cada LEAD_ID a su bloque de historial
        ws = libro["Estado actual"]
        _dar_formato(ws, 1, {
            "creacion_de_lead": FORMATO_FECHA,
            "FECHA_ULTIMO_MOVIMIENTO": FORMATO_FECHA,
        })
        for (celda,) in ws.iter_rows(min_row=2, max_col=1):
            if celda.value in inicio_bloques:
                celda.hyperlink = Hyperlink(
                    ref=celda.coordinate, location=f"'{HOJA_POR_LEAD}'!A{inicio_bloques[celda.value]}"
                )
                celda.font = Font(name=FUENTE, color="0563C1", underline="single")


# ---------------------------------------------------------------------------
# Función pública
# ---------------------------------------------------------------------------
def _carpeta_por_fecha(carpeta_base, fecha):
    """tablas + 23/09/2026 -> tablas/septiembre/23"""
    return Path(carpeta_base) / MESES[fecha.month - 1] / f"{fecha.day:02d}"

def _nombre_archivo(asesor):
    """'Miriam Gómez Cierres' -> 'reporte_estado_leads_miriam_gomez_cierres.xlsx'."""
    slug = re.sub(r"[^a-z0-9]+", "_", normalizar(asesor)).strip("_")
    return f"reporte_estado_leads_{slug}.xlsx" if slug else "reporte_estado_leads.xlsx"


def generar_reporte_estado_leads(historial, asesor="", ruta_salida=None,
                                 fecha_corte=None, carpeta_salida=None,
                                 verbose=True, fecha_archivo=None):
    """
    Genera el reporte de estado actual e historial de leads en Excel.

    Parámetros
    ----------
    historial : DataFrame o ruta (str/Path) a un .xlsx
        Resultado de la consulta: historial de etapas de los leads del asesor.
        Debe traer LEAD_ID, ETAPA_ANTERIOR, ETAPA_NUEVA y FECHA.
    asesor : str
        Nombre del asesor. Se muestra en el Resumen y se usa para nombrar el archivo.
    ruta_salida : str o Path, opcional
        Ruta completa del Excel. Si no se indica, se usa
        <carpeta_salida>/<mes>/<día>/reporte_estado_leads_<asesor>.xlsx.
    fecha_corte : datetime, opcional
        Fecha contra la que se calcula el tiempo transcurrido. Por defecto, ahora.
        Útil para que todos los reportes de una misma corrida usen la misma fecha.
    carpeta_salida : str o Path, opcional
        Carpeta donde se guarda el reporte cuando no se indica ruta_salida.
        Por defecto, tablas/reportes/individuales dentro del proyecto; el
        reporte queda en <carpeta_salida>/<mes>/<día>/.
    verbose : bool
        Si es True, imprime un resumen en consola.

    Devuelve
    --------
    Path del archivo generado, o None si la consulta no trajo registros.
    """
    if isinstance(historial, (str, Path)):
        historial = pd.read_excel(historial)

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in historial.columns]
    if faltantes:
        raise ValueError(f"A la consulta le faltan columnas requeridas: {faltantes}")

    if historial.empty:
        if verbose:
            print(f"[{asesor or 'sin asesor'}] La consulta no trajo registros; no se generó reporte.")
        return None

    fecha_corte = fecha_corte or datetime.now().replace(microsecond=0)
    if carpeta_salida is None:
        carpeta_salida = fn.ruta_proyecto(CARPETA_REPORTES)
    if ruta_salida:
        ruta_salida = Path(ruta_salida)
    else:
        fecha_archivo = fecha_archivo or fn.fecha_reportes()
        ruta_salida = _carpeta_por_fecha(carpeta_salida, fecha_archivo) / _nombre_archivo(asesor)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)

    hojas = generar_reporte(historial, fecha_corte)
    escribir_excel(hojas, ruta_salida, fecha_corte, asesor)

    if verbose:
        print(f"[{asesor or 'sin asesor'}] Reporte generado: {fn.ruta_para_mostrar(ruta_salida)} "
              f"({len(hojas['Estado actual'])} leads, {len(hojas['Historial'])} movimientos)")
    return ruta_salida


# ---------------------------------------------------------------------------
# Ejecución directa: ejemplo con un asesor
# ---------------------------------------------------------------------------
def main():
    labels = fn.import_json("json/secret.json")
    asesor = labels["people"]["sin_financiera_restringida"][0]  # primer asesor de la lista

    df = fn.import_xlsx(RUTA_HISTORIAL)
    historial = filtrar_por_etiqueta(df, asesor)
    generar_reporte_estado_leads(historial, asesor=asesor)


if __name__ == "__main__":
    main()