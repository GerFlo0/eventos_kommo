"""Genera ACUMULADOS.xlsx a partir de los archivos base .xlsx definidos en configuration.json.

Cada fuente produce una hoja del archivo de salida, aplicando en orden: filtro de ESTATUS,
exclusión de LEAD (opcional) y rango de fechas (inclusivo) sobre COLUMNA_FECHA. Las filas
se ordenan por esa misma fecha, conservando el orden del archivo base dentro de cada día.

Uso:
    python generar_acumulados.py
    python generar_acumulados.py --desde 01/09/2026 --hasta 28/09/2026
    python generar_acumulados.py --config otra_configuracion.json
"""

import argparse
import json
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from lectura import Celda, ErrorLectura, Tabla, leer_tabla, normalizar, serial_a_fecha

CONFIG_POR_DEFECTO = Path(__file__).resolve().parent / "json" / "configuration.json"
FORMATOS_FECHA_TEXTO = ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S")
MAX_EJEMPLOS_FILAS = 10

FUENTE_CAMPOS_REQUERIDOS = (
    "HOJA_DESTINO", "ARCHIVO", "HOJA_ORIGEN", "COLUMNA_FECHA",
    "COLUMNA_ESTATUS", "ESTATUS_INCLUIDOS", "COLUMNAS_SALIDA",
)


class ErrorConfiguracion(Exception):
    """Error con mensaje listo para mostrar al usuario."""


# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

def parsear_fecha(texto, campo: str) -> date | None:
    if texto is None or str(texto).strip() == "":
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(texto).strip(), formato).date()
        except ValueError:
            pass
    raise ErrorConfiguracion(f'{campo} "{texto}" no es válida. Usa DD/MM/AAAA o AAAA-MM-DD.')


def cargar_configuracion(ruta: Path) -> dict:
    if not ruta.exists():
        raise ErrorConfiguracion(f"No existe el archivo de configuración: {ruta}")
    try:
        config = json.loads(ruta.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ErrorConfiguracion(
            f"{ruta.name} no es un JSON válido: {e}\n"
            "Si usas rutas de Windows, escribe las diagonales dobles (C:\\\\...) o usa / (C:/...)."
        ) from e

    fuentes = config.get("fuentes") or config.get("FUENTES")
    if not fuentes:
        raise ErrorConfiguracion("configuration.json no tiene FUENTES definidas.")
    for i, fuente in enumerate(fuentes, start=1):
        faltantes = [c for c in FUENTE_CAMPOS_REQUERIDOS if not fuente.get(c)]
        if faltantes:
            raise ErrorConfiguracion(f"La fuente #{i} no tiene: {', '.join(faltantes)}")
    return config


def resolver_ruta(valor: str, base: Path) -> Path:
    ruta = Path(valor.strip()).expanduser()
    return ruta if ruta.is_absolute() else base / ruta


# ---------------------------------------------------------------------------
# Filtrado
# ---------------------------------------------------------------------------

def indice_columna(tabla: Tabla, nombre: str, hoja: str) -> int:
    buscado = normalizar(nombre)
    for i, encabezado in enumerate(tabla.encabezados):
        if normalizar(encabezado) == buscado:
            return i
    disponibles = ", ".join(f'"{e.strip()}"' for e in tabla.encabezados if e.strip())
    raise ErrorConfiguracion(f'No existe la columna "{nombre}" en la hoja "{hoja}". Columnas: {disponibles}')


def valor_a_fecha(valor) -> date | None:
    """Interpreta el valor de la columna de fecha; None si no es una fecha reconocible."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool) and 1 <= valor < 2958466:
        return serial_a_fecha(valor).date()  # celda con fecha pero sin formato de fecha
    if isinstance(valor, str) and valor.strip():
        for formato in FORMATOS_FECHA_TEXTO:
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def normalizar_lead(valor) -> str:
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return "" if valor is None else str(valor).strip()


def filtrar(tabla: Tabla, fuente: dict, desde: date, hasta: date | None) -> list[list[Celda]]:
    hoja = fuente["HOJA_DESTINO"]
    origen = fuente["HOJA_ORIGEN"]
    i_estatus = indice_columna(tabla, fuente["COLUMNA_ESTATUS"], origen)
    i_fecha = indice_columna(tabla, fuente["COLUMNA_FECHA"], origen)
    i_lead = indice_columna(tabla, fuente["COLUMNA_LEAD"], origen) if fuente.get("COLUMNA_LEAD") else None

    estatus_incluidos = {normalizar(e) for e in fuente["ESTATUS_INCLUIDOS"]}
    leads_excluidos = {normalizar_lead(l) for l in fuente.get("LEADS_EXCLUIDOS") or []}

    filas = tabla.filas
    resumen = [f"{len(filas)} leídas"]

    filas = [(n, c) for n, c in filas if normalizar(c[i_estatus].valor) in estatus_incluidos]
    resumen.append(f"{len(filas)} por estatus")

    if i_lead is not None and leads_excluidos:
        filas = [(n, c) for n, c in filas if normalizar_lead(c[i_lead].valor) not in leads_excluidos]
        resumen.append(f"{len(filas)} por lead")

    seleccionadas, sin_fecha = [], []
    for numero, celdas in filas:
        fecha = valor_a_fecha(celdas[i_fecha].valor)
        if fecha is None:
            sin_fecha.append(numero)
        elif fecha >= desde and (hasta is None or fecha <= hasta):
            seleccionadas.append((fecha, celdas))
    resumen.append(f"{len(seleccionadas)} por fecha")

    print(f"  [{hoja}] filas: " + " -> ".join(resumen))
    if sin_fecha:
        ejemplos = ", ".join(map(str, sin_fecha[:MAX_EJEMPLOS_FILAS]))
        extra = "..." if len(sin_fecha) > MAX_EJEMPLOS_FILAS else ""
        print(f'  [{hoja}] AVISO: {len(sin_fecha)} filas con "{fuente["COLUMNA_FECHA"]}" vacía o inválida '
              f"no se incluyeron (filas {ejemplos}{extra} del archivo base).")

    seleccionadas.sort(key=lambda par: par[0])  # sort estable: respeta el orden original en cada día
    return [celdas for _, celdas in seleccionadas]


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

def escribir_hoja(ws, tabla: Tabla, columnas_salida: list[str], filas: list[list[Celda]], hoja_origen: str):
    indices = [indice_columna(tabla, nombre, hoja_origen) for nombre in columnas_salida]
    fuente_normal = Font(name="Arial", size=10)
    fuente_encabezado = Font(name="Arial", size=10, bold=True)

    for j, i in enumerate(indices, start=1):
        celda = ws.cell(row=1, column=j, value=tabla.encabezados[i])
        celda.font = fuente_encabezado

    for r, celdas in enumerate(filas, start=2):
        for j, i in enumerate(indices, start=1):
            origen = celdas[i]
            celda = ws.cell(row=r, column=j, value=origen.valor)
            celda.font = fuente_normal
            if origen.formato and origen.formato != "General":
                celda.number_format = origen.formato
            elif isinstance(origen.valor, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm:ss"
            elif isinstance(origen.valor, date):
                celda.number_format = "dd/mm/yyyy"
            elif isinstance(origen.valor, time):
                celda.number_format = "hh:mm:ss"

    for j, i in enumerate(indices, start=1):
        ancho = max([len(tabla.encabezados[i].strip())] +
                    [len(str(f[i].valor)) for f in filas[:500] if f[i].valor is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 10), 45)

    ws.freeze_panes = "A2"
    if indices:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(indices))}{len(filas) + 1}"


# ---------------------------------------------------------------------------

def ejecutar(ruta_config: Path, desde_cli: str | None, hasta_cli: str | None):
    config = cargar_configuracion(ruta_config)
    settings = config.get("settings") or {}
    fuentes = config.get("fuentes") or config["FUENTES"]
    base = ruta_config.parent.parent if ruta_config.parent.name == "json" else ruta_config.parent

    desde = parsear_fecha(desde_cli or settings.get("FECHA_DESDE") or config.get("FECHA_DESDE"), "FECHA_DESDE")
    hasta = parsear_fecha(hasta_cli or settings.get("FECHA_HASTA") or config.get("FECHA_HASTA"), "FECHA_HASTA")
    if desde is None:
        raise ErrorConfiguracion("Indica FECHA_DESDE en configuration.json o con --desde DD/MM/AAAA.")
    if hasta is not None and hasta < desde:
        raise ErrorConfiguracion(f"FECHA_HASTA ({hasta:%d/%m/%Y}) es anterior a FECHA_DESDE ({desde:%d/%m/%Y}).")

    salida = resolver_ruta(settings.get("RUTA_ACUMULADOS") or config.get("ARCHIVO_SALIDA") or "ACUMULADOS.xlsx", base)

    rango = f"desde {desde:%d/%m/%Y}" + (f" hasta {hasta:%d/%m/%Y}" if hasta else " (sin fecha límite)")
    print(f"Generando acumulados {rango}")

    libro = Workbook()
    libro.remove(libro.active)
    for fuente in fuentes:
        destino = fuente["HOJA_DESTINO"]
        archivo = resolver_ruta(fuente["ARCHIVO"], base)
        print(f'- Leyendo "{archivo.name}" / "{fuente["HOJA_ORIGEN"]}" para "{destino}"...')
        try:
            tabla = leer_tabla(archivo, fuente["HOJA_ORIGEN"])
            filas = filtrar(tabla, fuente, desde, hasta)
            ws = libro.create_sheet(title=destino[:31])
            escribir_hoja(ws, tabla, fuente["COLUMNAS_SALIDA"], filas, fuente["HOJA_ORIGEN"])
        except (ErrorLectura, ErrorConfiguracion) as e:
            raise ErrorConfiguracion(f"[{destino}] {e}") from e

    salida.parent.mkdir(parents=True, exist_ok=True)
    try:
        libro.save(salida)
    except PermissionError as e:
        raise ErrorConfiguracion(f"No se pudo guardar {salida}. ¿Está abierto en Excel? Ciérralo e intenta de nuevo.") from e
    print(f"Listo: {salida}")


def main():
    parser = argparse.ArgumentParser(description="Genera ACUMULADOS.xlsx a partir de los archivos base.")
    parser.add_argument("--desde", help="Fecha inicial DD/MM/AAAA (sobrescribe FECHA_DESDE)")
    parser.add_argument("--hasta", help="Fecha final DD/MM/AAAA (sobrescribe FECHA_HASTA)")
    parser.add_argument("--config", type=Path, default=CONFIG_POR_DEFECTO, help="Ruta de configuration.json")
    args = parser.parse_args()

    try:
        ejecutar(args.config.resolve(), args.desde, args.hasta)
    except ErrorConfiguracion as e:
        print(f"\nERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
