"""Lectura de una hoja de un archivo .xlsx como tabla (encabezados + filas con formato)."""

import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

EPOCA_EXCEL = datetime(1899, 12, 30)  # día 0 de los números de serie de Excel


class ErrorLectura(Exception):
    """Error con mensaje listo para mostrar al usuario."""


@dataclass
class Celda:
    valor: object = None
    formato: str | None = None


@dataclass
class Tabla:
    encabezados: list[str]
    filas: list[tuple[int, list[Celda]]] = field(default_factory=list)  # (número de fila, celdas)


def normalizar(texto) -> str:
    """Mayúsculas, sin acentos y sin espacios sobrantes (para comparar encabezados/valores)."""
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.split()).upper()


def serial_a_fecha(serial: float) -> datetime:
    """Número de serie de Excel (días desde 1899-12-30) a datetime."""
    return EPOCA_EXCEL + timedelta(seconds=round(serial * 86400))


def leer_tabla(ruta: Path, hoja: str) -> Tabla:
    """Lee la hoja indicada usando los valores calculados (no las fórmulas)."""
    import openpyxl

    if not ruta.exists():
        raise ErrorLectura(f"No existe el archivo: {ruta}")

    # openpyxl avisa que no soporta algunas validaciones de datos; no afecta la lectura
    warnings.filterwarnings("ignore", message="Data Validation extension", category=UserWarning)
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as e:
        raise ErrorLectura(f'No se pudo abrir "{ruta.name}" como Excel: {e}') from e
    try:
        ws = libro[_buscar_hoja(hoja, libro.sheetnames, ruta.name)]
        ws.reset_dimensions()  # evita dimensiones mal guardadas que recorten filas
        filas = [
            (numero, [Celda(c.value, getattr(c, "number_format", None)) for c in fila])
            for numero, fila in enumerate(ws.iter_rows(), start=1)
        ]
    finally:
        libro.close()
    return _construir_tabla(filas, hoja)


def _buscar_hoja(buscada: str, disponibles: list[str], libro: str) -> str:
    if buscada in disponibles:
        return buscada
    for nombre in disponibles:
        if normalizar(nombre) == normalizar(buscada):
            return nombre
    raise ErrorLectura(
        f'No existe la hoja "{buscada}" en "{libro}". Hojas disponibles: {", ".join(disponibles)}'
    )


def _construir_tabla(filas: list[tuple[int, list[Celda]]], hoja: str) -> Tabla:
    """Toma la primera fila no vacía como encabezados y descarta filas totalmente vacías."""
    filas = [(n, celdas) for n, celdas in filas if any(c.valor not in (None, "") for c in celdas)]
    if not filas:
        raise ErrorLectura(f'La hoja "{hoja}" está vacía.')

    _, fila_encabezados = filas[0]
    encabezados = ["" if c.valor is None else str(c.valor) for c in fila_encabezados]
    while encabezados and encabezados[-1].strip() == "":
        encabezados.pop()

    ancho = len(encabezados)
    datos = []
    for numero, celdas in filas[1:]:
        celdas = celdas[:ancho] + [Celda() for _ in range(ancho - len(celdas))]
        datos.append((numero, celdas))
    return Tabla(encabezados=encabezados, filas=datos)
