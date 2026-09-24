"""
Reporte general de efectividad de todos los asesores.

Lee la hoja "Estado actual" de cada reporte individual de un día y genera un
Excel con dos hojas:
  - Efectividad: conteo por etapa y % de efectividad neta y bruta por asesor.
  - Detalle completo: cada lead con su asesor, etapa actual y categoría.
    Los conteos de "Efectividad" son fórmulas COUNTIFS sobre esta hoja.

Rutas por defecto (según la fecha de ejecución, p. ej. 23 de septiembre),
dentro de la carpeta del proyecto:
  - Reportes individuales: tablas/reportes/individuales/septiembre/23
  - Reporte general:       tablas/reportes/generales/septiembre/23

Uso:
    python general_report.py                      # reportes de hoy
    python general_report.py --fecha 22/09/2026   # reportes de otro día
    python general_report.py --fecha 22/09/2026 --descripcion "SIN FINANCIERA RESTRINGIDA"

Desde otro script:
    from general_report import generar_reporte_general
    generar_reporte_general()                          # hoy
    generar_reporte_general(fecha=date(2026, 9, 22))   # otro día
"""

import argparse
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

import functions as fn
from individual_reports import MESES, clasificar_etapa, normalizar

CARPETA_INDIVIDUALES = "tablas/reportes/individuales"
CARPETA_GENERALES = "tablas/reportes/generales"
PREFIJO_INDIVIDUAL = "reporte_estado_leads"
HOJA_ESTADO = "Estado actual"

# Categoría de individual_reports.py -> columna del reporte general
COLUMNA_POR_CATEGORIA = {
    "Logrado con éxito": "Ganados",
    "Oferta": "Oferta",
    "Oferta en espera": "Oferta en Espera",
    "Documentación": "Documentación",
    "Capturado": "Capturado",
    "Venta perdida": "Lead Perdido",
    "Sin capacidad": "Sin Capacidad",
    "Dejado a futuro": "Bimestre",
    "Otros": "Otros",
}
# Etapas que forman los "Negocios cerrables"
CERRABLES = ["Ganados", "Oferta", "Oferta en Espera", "Documentación", "Capturado", "Lead Perdido"]
# Se suman a los cerrables para obtener el "Total asignación"
FUERA_DE_CERRABLES = ["Sin Capacidad", "Bimestre"]

# Color de relleno por columna cuando el valor es mayor a 0
COLORES = {
    "Ganados": "C8E6C9",
    "Oferta": "FFF9C4",
    "Oferta en Espera": "FFE0B2",
    "Documentación": "FFF3CD",
    "Capturado": "FFE9A8",
    "Lead Perdido": "FFCDD2",
    "Sin Capacidad": "FFD9B3",
    "Otros": "E0E0E0",
}

FUENTE = "Arial"
AZUL = "1F4E79"
# % Efect. BRUTA a partir del cual la celda se pinta de verde
UMBRAL_BRUTA = 0.20
VERDE = "C8E6C9"
FORMATO_FECHA = "dd/mm/yyyy hh:mm"
_RELLENO_AZUL = PatternFill("solid", start_color=AZUL)
_BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Lectura de reportes individuales
# ---------------------------------------------------------------------------
def carpeta_del_dia(carpeta_base, fecha):
    """tablas/reportes/generales + 23/09/2026 -> tablas/reportes/generales/septiembre/23"""
    return Path(carpeta_base) / MESES[fecha.month - 1] / f"{fecha.day:02d}"


def _nombre_asesor(ruta):
    """Toma el asesor de la celda B2 del Resumen; si no está, lo deduce del nombre del archivo."""
    try:
        wb = load_workbook(ruta, read_only=True)
        try:
            if "Resumen" in wb.sheetnames:
                valor = wb["Resumen"]["B2"].value
                if valor:
                    return str(valor).strip()
        finally:
            wb.close()   # en modo read_only el archivo queda abierto si no se cierra
    except Exception:
        pass
    nombre = ruta.stem.replace(PREFIJO_INDIVIDUAL, "").strip("_")
    return nombre.replace("_", " ").title() or ruta.stem


def leer_reportes_individuales(carpeta, verbose=True):
    """Une la hoja 'Estado actual' de todos los reportes individuales de la carpeta."""
    carpeta = Path(carpeta)
    if not carpeta.is_dir():
        raise FileNotFoundError(f"No existe la carpeta de reportes individuales: {carpeta}")

    archivos = sorted(
        p for p in carpeta.glob(f"{PREFIJO_INDIVIDUAL}*.xlsx") if not p.name.startswith("~$")
    )
    if not archivos:
        raise FileNotFoundError(f"No hay reportes individuales en: {carpeta}")

    tablas = []
    for ruta in archivos:
        try:
            df = pd.read_excel(ruta, sheet_name=HOJA_ESTADO)
        except ValueError:
            if verbose:
                print(f"  Omitido (sin hoja '{HOJA_ESTADO}'): {ruta.name}")
            continue
        df.insert(0, "ASESOR", _nombre_asesor(ruta))
        tablas.append(df)
        if verbose:
            print(f"  {ruta.name}: {len(df)} leads")

    if not tablas:
        raise ValueError(f"Ningún archivo de {carpeta} tiene la hoja '{HOJA_ESTADO}'.")

    datos = pd.concat(tablas, ignore_index=True)
    datos["CATEGORIA"] = datos["ETAPA_ACTUAL"].map(clasificar_etapa).map(COLUMNA_POR_CATEGORIA)
    return datos


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _titulo(ws, texto, n_cols):
    ws.cell(row=1, column=1, value=texto)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    for col in range(1, n_cols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = _RELLENO_AZUL
        c.font = Font(name=FUENTE, bold=True, size=12, color="FFFFFF")
        c.alignment = _CENTRO
    ws.row_dimensions[1].height = 25.5


def _encabezados(ws, nombres, fila=2):
    for col, nombre in enumerate(nombres, start=1):
        c = ws.cell(row=fila, column=col, value=nombre)
        c.fill = _RELLENO_AZUL
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.alignment = _CENTRO
        c.border = _BORDE
    ws.row_dimensions[fila].height = 45.75


def _hoja_detalle(wb, datos, orden_asesores, titulo):
    ws = wb.create_sheet("Detalle completo")
    columnas = [
        ("Nombre", "ASESOR", 28, None),
        ("ID Lead", "LEAD_ID", 12, None),
        ("Etapa actual (Kommo)", "ETAPA_ACTUAL", 22, None),
        ("Categoría", "CATEGORIA", 16, None),
        ("¿Ganado?", None, 10, None),
        ("Fecha último movimiento", "FECHA_ULTIMO_MOVIMIENTO", 20, FORMATO_FECHA),
        ("Tiempo transcurrido", "TIEMPO_TRANSCURRIDO", 16, None),
        ("Total movimientos", "TOTAL_MOVIMIENTOS", 12, None),
        ("Ruta", "RUTA", 90, None),
    ]
    _titulo(ws, titulo, len(columnas))
    _encabezados(ws, [c[0] for c in columnas])

    datos = datos.copy()
    datos["_orden"] = datos["ASESOR"].map({a: i for i, a in enumerate(orden_asesores)})
    datos = datos.sort_values(["_orden", "LEAD_ID"])

    fila = 3
    for _, lead in datos.iterrows():
        for col, (_, campo, _, fmt) in enumerate(columnas, start=1):
            if campo is None:
                valor = "SÍ" if lead["CATEGORIA"] == "Ganados" else None
            else:
                valor = lead.get(campo)
                if pd.isna(valor):
                    valor = None
                elif isinstance(valor, pd.Timestamp):
                    valor = valor.to_pydatetime()
            c = ws.cell(row=fila, column=col, value=valor)
            c.font = Font(name=FUENTE, bold=(col == 1))
            c.border = _BORDE
            if fmt:
                c.number_format = fmt
            if col not in (1, 3, 9):
                c.alignment = Alignment(horizontal="center")
        fila += 1

    for col, (_, _, ancho, _) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columnas))}{fila - 1}"
    return ws, fila - 1


def _hoja_efectividad(wb, datos, orden_asesores, titulo, ultima_fila_detalle):
    ws = wb.active
    ws.title = "Efectividad"

    hay_otros = (datos["CATEGORIA"] == "Otros").any()
    extras = FUERA_DE_CERRABLES + (["Otros"] if hay_otros else [])
    columnas = ["Nombre", "Negocios cerrables"] + CERRABLES + extras + [
        "Total asignación", "% Efect. NETA", "% Efect. BRUTA"
    ]
    letra = {nombre: get_column_letter(i) for i, nombre in enumerate(columnas, start=1)}

    _titulo(ws, titulo, len(columnas))
    _encabezados(ws, columnas)

    rango_nombre = f"'Detalle completo'!$A$3:$A${ultima_fila_detalle}"
    rango_categoria = f"'Detalle completo'!$D$3:$D${ultima_fila_detalle}"
    conteos = datos.groupby(["ASESOR", "CATEGORIA"]).size()

    fila = 3
    for asesor in orden_asesores:
        ws.cell(row=fila, column=1, value=asesor).font = Font(name=FUENTE, bold=True)

        for cat in CERRABLES + extras:
            col = columnas.index(cat) + 1
            c = ws.cell(
                row=fila, column=col,
                value=f'=COUNTIFS({rango_nombre},$A{fila},{rango_categoria},"{cat}")',
            )
            c.font = Font(name=FUENTE, bold=cat in ("Ganados", "Bimestre"))
            if conteos.get((asesor, cat), 0) > 0 and cat in COLORES:
                c.fill = PatternFill("solid", start_color=COLORES[cat])

        L = letra
        ws[f"{L['Negocios cerrables']}{fila}"] = f"=SUM({L[CERRABLES[0]]}{fila}:{L[CERRABLES[-1]]}{fila})"
        ws[f"{L['Total asignación']}{fila}"] = "=" + "+".join(
            f"{L[c]}{fila}" for c in ["Negocios cerrables"] + extras
        )
        ws[f"{L['% Efect. NETA']}{fila}"] = f"=IFERROR({L['Ganados']}{fila}/{L['Negocios cerrables']}{fila},0)"
        ws[f"{L['% Efect. BRUTA']}{fila}"] = f"=IFERROR({L['Ganados']}{fila}/{L['Total asignación']}{fila},0)"

        ws[f"{L['Negocios cerrables']}{fila}"].font = Font(name=FUENTE)
        ws[f"{L['Total asignación']}{fila}"].font = Font(name=FUENTE, bold=True)
        ws[f"{L['% Efect. NETA']}{fila}"].font = Font(name=FUENTE)
        ws[f"{L['% Efect. BRUTA']}{fila}"].font = Font(name=FUENTE, bold=True)
        ws[f"{L['% Efect. BRUTA']}{fila}"].fill = PatternFill("solid", start_color="FFCDD2")
        fila += 1

    # Fila TOTAL
    primera, ultima = 3, fila - 1
    ws.cell(row=fila, column=1, value="TOTAL")
    for nombre in columnas[1:]:
        col = letra[nombre]
        if nombre == "% Efect. NETA":
            valor = f"=IFERROR({letra['Ganados']}{fila}/{letra['Negocios cerrables']}{fila},0)"
        elif nombre == "% Efect. BRUTA":
            valor = f"=IFERROR({letra['Ganados']}{fila}/{letra['Total asignación']}{fila},0)"
        else:
            valor = f"=SUM({col}{primera}:{col}{ultima})"
        ws[f"{col}{fila}"] = valor
    for col in range(1, len(columnas) + 1):
        c = ws.cell(row=fila, column=col)
        c.fill = _RELLENO_AZUL
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")

    # Formato general
    for r in range(3, fila + 1):
        for col in range(1, len(columnas) + 1):
            c = ws.cell(row=r, column=col)
            c.border = _BORDE
            if col > 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
        for nombre in ("% Efect. NETA", "% Efect. BRUTA"):
            ws[f"{letra[nombre]}{r}"].number_format = "0.0%"

    
        # % Efect. BRUTA en verde cuando es mayor o igual al umbral (solo filas de asesores)
    col_bruta = letra["% Efect. BRUTA"]
    ws.conditional_formatting.add(
        f"{col_bruta}{primera}:{col_bruta}{ultima}",
        CellIsRule(operator="greaterThanOrEqual", formula=[str(UMBRAL_BRUTA)],
                   fill=PatternFill("solid", start_color=VERDE, end_color=VERDE)),
    )
    ws.column_dimensions["A"].width = 30
    for nombre in columnas[1:]:
        ws.column_dimensions[letra[nombre]].width = 12
    ws.freeze_panes = "B3"
    return ws


def _orden_por_efectividad(datos):
    """Asesores ordenados por % Efect. NETA (desc), luego por ganados y nombre."""
    resumen = []
    for asesor, grupo in datos.groupby("ASESOR"):
        cerrables = grupo["CATEGORIA"].isin(CERRABLES).sum()
        ganados = (grupo["CATEGORIA"] == "Ganados").sum()
        resumen.append((asesor, ganados / cerrables if cerrables else 0, ganados))
    resumen.sort(key=lambda x: (-x[1], -x[2], normalizar(x[0])))
    return [a for a, _, _ in resumen]


# ---------------------------------------------------------------------------
# Función pública
# ---------------------------------------------------------------------------
def generar_reporte_general(fecha=None, carpeta_individuales=None, carpeta_salida=None,
                            ruta_salida=None, descripcion="", verbose=True):
    """
    Genera el reporte general de efectividad.

    Parámetros
    ----------
    fecha : date o datetime, opcional
        Día de los reportes a consolidar. Por defecto, hoy.
    carpeta_individuales : str o Path, opcional
        Carpeta con los reportes individuales. Por defecto,
        tablas/reportes/individuales/<mes>/<día> según `fecha`.
    carpeta_salida : str o Path, opcional
        Carpeta del reporte general. Por defecto,
        tablas/reportes/generales/<mes>/<día> según `fecha`.
    ruta_salida : str o Path, opcional
        Ruta completa del Excel; si se da, ignora carpeta_salida.
    descripcion : str
        Texto opcional para el título (p. ej. "SIN FINANCIERA RESTRINGIDA").

    Devuelve
    --------
    Path del archivo generado.
    """
    fecha = fecha or date.today()
    carpeta_individuales = Path(carpeta_individuales or
                                carpeta_del_dia(fn.ruta_proyecto(CARPETA_INDIVIDUALES), fecha))
    if ruta_salida:
        ruta_salida = Path(ruta_salida)
    else:
        carpeta_salida = Path(carpeta_salida or
                              carpeta_del_dia(fn.ruta_proyecto(CARPETA_GENERALES), fecha))
        ruta_salida = carpeta_salida / f"reporte_general_{fecha:%Y-%m-%d}.xlsx"

    if verbose:
        print(f"Leyendo reportes individuales de: {fn.ruta_para_mostrar(carpeta_individuales)}")
    datos = leer_reportes_individuales(carpeta_individuales, verbose)
    orden = _orden_por_efectividad(datos)

    partes = [f"{MESES[fecha.month - 1].upper()} {fecha.year}", "REPORTE GENERAL DE EFECTIVIDAD"]
    if descripcion:
        partes.append(descripcion.upper())
    partes.append(f"corte {fecha:%d/%m/%Y}")
    base_titulo = " — ".join(partes)

    wb = Workbook()
    _, ultima_fila = _hoja_detalle(wb, datos, orden, f"{base_titulo} — DETALLE COMPLETO")
    _hoja_efectividad(wb, datos, orden, f"{base_titulo} · ordenado por % Efect. NETA", ultima_fila)

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)

    if verbose:
        otros = datos.loc[datos["CATEGORIA"] == "Otros", "ETAPA_ACTUAL"].unique()
        print(f"Reporte general generado: {fn.ruta_para_mostrar(ruta_salida)} "
              f"({len(orden)} asesores, {len(datos)} leads)")
        if len(otros):
            print(f"  Aviso: etapas sin clasificar (columna 'Otros'): {', '.join(map(str, otros))}")
        # Un lead con etiquetas de dos asesores sale en ambos reportes
        # individuales y se cuenta en los dos.
        por_lead = datos.groupby("LEAD_ID")["ASESOR"].nunique()
        repetidos = por_lead[por_lead > 1].index.tolist()
        if repetidos:
            print(f"  Aviso: {len(repetidos)} lead(s) aparecen con más de un asesor y se "
                  f"cuentan en cada uno: {', '.join(map(str, repetidos))}")
    return ruta_salida


# ---------------------------------------------------------------------------
def _leer_fecha(texto):
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise argparse.ArgumentTypeError(f"Fecha no válida: '{texto}'. Usa dd/mm/aaaa o aaaa-mm-dd.")


def main():
    parser = argparse.ArgumentParser(description="Reporte general de efectividad de asesores.")
    parser.add_argument("--fecha", type=_leer_fecha,
                        help="Día de los reportes a consolidar (dd/mm/aaaa o aaaa-mm-dd). Por defecto, hoy.")
    parser.add_argument("--carpeta", help="Carpeta de reportes individuales (sustituye la ruta por fecha).")
    parser.add_argument("--salida", help="Carpeta del reporte general (sustituye la ruta por fecha).")
    parser.add_argument("--descripcion", default="", help="Texto adicional para el título.")
    args = parser.parse_args()

    try:
        generar_reporte_general(
            fecha=args.fecha,
            carpeta_individuales=args.carpeta,
            carpeta_salida=args.salida,
            descripcion=args.descripcion,
        )
    except (FileNotFoundError, ValueError) as error:
        raise SystemExit(f"No se generó el reporte general: {error}")


if __name__ == "__main__":
    main()