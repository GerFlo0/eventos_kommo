import contextlib
import io
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import load_workbook

from _helpers import ASESORES, SECRET, cargar_extractor, fake_kommo, fn

import app
import general_report as gr
import individual_reports as ir

FECHA_CORTE = datetime(2026, 9, 23, 12, 0, 0)


def generar_historial(carpeta):
    """Corre el extractor contra la API simulada y devuelve el historial."""
    ex = cargar_extractor()
    ex.SALIDA = str(Path(carpeta) / "historial.xlsx")
    fake = fake_kommo.FakeKommo()
    with mock.patch.object(ex.SESSION, "get", fake.get), mock.patch.object(ex.time, "sleep"), \
            contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        ex.main()
    return fn.import_xlsx(ex.SALIDA)


class TestClasificacion(unittest.TestCase):
    def test_categorias(self):
        casos = {
            "Logrado con éxito": "Logrado con éxito",
            "LOGRADO CON EXITO": "Logrado con éxito",
            "Oferta": "Oferta",
            "Oferta en espera": "Oferta en espera",
            "Documentación ": "Documentación",
            "Capturado": "Capturado",
            "Ventas perdidas": "Venta perdida",
            "SIN CAPACIDAD": "Sin capacidad",
            "BIMESTRE MAR-ABR 2026": "Dejado a futuro",
            "Enero 2027": "Dejado a futuro",
            "2025": "Dejado a futuro",
            "OFERTA 20%": "Otros",
            "Llamar en 20 días": "Otros",
            "Código 12026": "Otros",
            "Contactado": "Otros",
            None: "Otros",
        }
        for etapa, esperado in casos.items():
            with self.subTest(etapa=etapa):
                self.assertEqual(ir.clasificar_etapa(etapa), esperado)

    def test_formato_tiempo(self):
        self.assertEqual(ir.formato_tiempo(12.4), "12 d 10 h")
        self.assertEqual(ir.formato_tiempo(None), "")


class TestReportes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.historial = generar_historial(cls.dir)
        cls.individuales = cls.dir / "individuales"
        cls.rutas = {}
        with contextlib.redirect_stdout(io.StringIO()):
            for asesor in ASESORES:
                datos = ir.filtrar_por_etiqueta(cls.historial, asesor)
                cls.rutas[asesor] = ir.generar_reporte_estado_leads(
                    datos, asesor=asesor, fecha_corte=FECHA_CORTE,
                    carpeta_salida=cls.individuales)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_reporte_individual(self):
        ruta = self.rutas["ASESOR UNO CIERRES"]
        self.assertEqual(ruta.parent, self.individuales / "septiembre" / "23")
        self.assertEqual(ruta.name, "reporte_estado_leads_asesor_uno_cierres.xlsx")
        wb = load_workbook(ruta)
        self.assertEqual(wb.sheetnames, ["Historial por lead", "Resumen", "Estado actual"])
        self.assertEqual(wb["Resumen"]["B2"].value, "ASESOR UNO CIERRES")
        estado = wb["Estado actual"]
        self.assertTrue(all(c.hyperlink for (c,) in estado.iter_rows(min_row=2, max_col=1)))
        notas = [c.value for (c,) in wb["Resumen"].iter_rows(max_col=1) if c.value]
        self.assertTrue(any("contiene un año" in str(n) for n in notas))

    def test_sin_registros_no_genera_archivo(self):
        vacio = self.historial.iloc[0:0]
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertIsNone(ir.generar_reporte_estado_leads(vacio, asesor="X"))

    def test_faltan_columnas(self):
        with self.assertRaises(ValueError):
            ir.generar_reporte_estado_leads(pd.DataFrame({"LEAD_ID": [1]}))

    def test_carpeta_por_defecto_dentro_del_proyecto(self):
        datos = ir.filtrar_por_etiqueta(self.historial, "ASESOR UNO CIERRES")
        destino = self.dir / "por_defecto"
        with mock.patch.object(ir.fn, "ruta_proyecto", return_value=destino) as rp, \
                contextlib.redirect_stdout(io.StringIO()):
            ruta = ir.generar_reporte_estado_leads(datos, asesor="A B", fecha_corte=FECHA_CORTE)
        rp.assert_called_once_with(ir.CARPETA_REPORTES)
        self.assertTrue(ruta.exists())
        self.assertEqual(ruta.parent, destino / "septiembre" / "23")

    def test_reporte_general(self):
        salida = io.StringIO()
        with contextlib.redirect_stdout(salida):
            ruta = gr.generar_reporte_general(
                fecha=date(2026, 9, 23),
                carpeta_individuales=self.individuales / "septiembre" / "23",
                ruta_salida=self.dir / "general.xlsx")
        wb = load_workbook(ruta)
        self.assertEqual(wb.sheetnames, ["Efectividad", "Detalle completo"])
        ef = wb["Efectividad"]
        self.assertTrue(str(ef["C3"].value).startswith("=COUNTIFS("))
        self.assertEqual(ef.cell(row=ef.max_row, column=1).value, "TOTAL")
        # "ANA PRUEBA" coincide dentro de "JULIANA PRUEBA": se avisa del doble conteo
        self.assertIn("más de un asesor", salida.getvalue())
        # "Seguimiento 20 días" ya no es "Dejado a futuro"
        self.assertIn("Seguimiento 20 días", salida.getvalue())

    def test_nombre_asesor_cierra_el_archivo(self):
        libro = mock.MagicMock()
        libro.sheetnames = ["Resumen"]
        libro.__getitem__.return_value.__getitem__.return_value.value = "Nombre X"
        with mock.patch.object(gr, "load_workbook", return_value=libro):
            self.assertEqual(gr._nombre_asesor(Path("reporte_estado_leads_x.xlsx")), "Nombre X")
        libro.close.assert_called_once()

    def test_nombre_asesor_por_nombre_de_archivo(self):
        with mock.patch.object(gr, "load_workbook", side_effect=OSError):
            self.assertEqual(gr._nombre_asesor(Path("reporte_estado_leads_ana_prueba.xlsx")),
                             "Ana Prueba")


class TestApp(unittest.TestCase):
    def test_genera_un_reporte_por_asesor_con_datos(self):
        with tempfile.TemporaryDirectory() as tmp:
            historial = generar_historial(tmp)
            original = ir.generar_reporte_estado_leads

            def en_tmp(*args, **kwargs):
                return original(*args, carpeta_salida=Path(tmp) / "rep", **kwargs)

            secret = dict(SECRET)
            secret["asesores"] = ASESORES + ["NADIE"]
            with mock.patch.object(app.fn, "import_json", return_value=secret), \
                    mock.patch.object(app.fn, "import_xlsx", return_value=historial), \
                    mock.patch.object(app, "generar_reporte_estado_leads", side_effect=en_tmp), \
                    contextlib.redirect_stdout(io.StringIO()):
                app.main()
            archivos = sorted(p.name for p in (Path(tmp) / "rep").rglob("*.xlsx"))
        self.assertEqual(archivos, [
            "reporte_estado_leads_ana_prueba_cierres.xlsx",
            "reporte_estado_leads_asesor_uno_cierres.xlsx",
            "reporte_estado_leads_juliana_prueba_cierres.xlsx",
        ])


if __name__ == "__main__":
    unittest.main()
